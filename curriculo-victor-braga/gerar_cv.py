"""Gera o CV (DOCX) de Victor Braga em formato ATS-friendly: uma coluna, sem tabelas,
sem caixas de texto, cabeçalhos padrão. Depois converta para PDF com LibreOffice:
    soffice --headless --convert-to pdf Victor_Braga_CV.docx
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = "Calibri"
DARK = RGBColor(0x1F, 0x1F, 0x1F)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4 (padrão na Austrália)
sec.top_margin = sec.bottom_margin = Cm(1.3)
sec.left_margin = sec.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal.font.color.rgb = DARK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.08


def para(text="", size=10, bold=False, color=None, align=None, after=0, before=0, italic=False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F4E79"}.items():
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def heading(text):
    p = para(text.upper(), size=11, bold=True, color=ACCENT, before=8, after=3)
    bottom_border(p)


def job(title, company, dates):
    p = para(before=4)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    p.add_run(f" | {company} | {dates}").font.size = Pt(10)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(it)
        r.font.size = Pt(10)


def line(label, text):
    p = para(after=1)
    p.add_run(label + ": ").bold = True
    p.add_run(text)


# ---------- Cabeçalho ----------
para("VICTOR BRAGA", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Graphic Designer | Front of House | Kitchen Hand | Brand Ambassador", size=12, color=ACCENT,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Fortitude Valley, Brisbane QLD | 0451 084 505 | victorbcsa@gmail.com",
     size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Portfolio: [ADD PORTFOLIO LINK]", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

# ---------- Resumo ----------
heading("Profile")
para("Graphic designer with 7 years of agency experience in Brazil, producing print and point-of-sale "
     "material, large-format signage and digital artwork for brands such as Cinemark, "
     "Wickbold and Seven Boys. Looking for a role in screen printing, print production, "
     "signage or a creative studio.", after=3)
para("Now based in Brisbane, with 2026 experience in Australian hospitality and promotions. "
     "Worked front of house, kitchen and pizza chef roles, and as a brand ambassador for Activ8 "
     "at Woolworths activations and Riverfire. Trusted to open and close venues, manage stock "
     "and lead on-site activations.", after=2)

# ---------- Experiência ----------
heading("Australian Work Experience")
job("Front of House", "Austino Cuisine – Maroochydore Club, Sunshine Coast", "2026")
bullets([
    "Served food and drinks, took orders accurately and kept the floor clean and presentable.",
    "Organised food distribution from the pass and supported the kitchen and dishwashing station when needed.",
])

job("Front of House / Kitchen Hand / Pizza Chef", "Baretto Pizza, Brisbane", "2026")
bullets([
    "Opened and closed the restaurant and kitchen, including all cleaning procedures.",
    "Managed stock, ordering and storage, plus Uber Eats and takeaway orders.",
    "Prepared ingredients and cooked pizzas during busy services, following food safety standards.",
])

job("Brand Ambassador & Promotional Staff", "Activ8, Brisbane", "2026")
bullets([
    "Woolworths – Product Sampling & On-site Activation Lead: engaged customers, introduced "
    "products and handed out samples. Led on site, managing stock and deliveries, making "
    "operational decisions and reporting to campaign coordinators.",
    "Australian Retirement Trust – Mascot Minder, Riverfire 2026 (Brisbane Festival): "
    "accompanied the mascot Artie, supported interactions with the public and guided it "
    "safely through large crowds.",
    "Street & Event Promotion: handed out flyers in busy pedestrian areas and explained event "
    "details to the public.",
])

heading("Design Experience")
job("Senior Art Director / Graphic Designer", "Essenza Comunicação, São Paulo",
    "Dec 2020 – Feb 2026")
bullets([
    "Designed key visuals, in-store and point-of-sale material, large-format outdoor ads and "
    "social content for Cinemark (Brazil and Argentina), Wickbold and Seven Boys.",
    "Lead designer on brand collaborations, incl. Seven Boys x Sonic, Seven Boys x GOAT and "
    "Cinemark with Coca-Cola, Monster Energy and Free Fire.",
    "Created and named the Cinemark Club loyalty program, featured as a major sponsor at CCXP 2022.",
    "Adapted each campaign into multiple print and digital formats, often on tight deadlines.",
])

job("Graphic Designer", "Do It Comunicação, Petrópolis", "Mar 2019 – Sep 2020")
bullets([
    "Designed and prepared print files for flyers, posters, menus, labels and signage for "
    "supermarkets, craft breweries (Brewpoint, Odin, Da Corte) and restaurants.",
])

job("Photographer (Freelance)", "Petrópolis", "Feb 2017 – Mar 2019")
bullets([
    "Product, food and event photography for local businesses, incl. lighting, shooting and retouching.",
])

# ---------- Skills ----------
heading("Skills")
para("Meeting deadlines, attention to detail, teamwork, customer service, stock control, "
     "hands-on and physical work", after=1)

# ---------- Educação ----------
heading("Education")
para("Graphit School of Design, Petrópolis, Brazil (2017 – 2018): Graphic Design, Advertising Design, "
     "Web Design, Digital Illustration, Photo Retouching", after=1)
para("Exchange program, Lordstown High School, Ohio, USA (2015)", after=1)

# ---------- Idiomas ----------
heading("Languages")
para("English (fluent) | Portuguese (native) | Spanish (basic)")

doc.save("Victor_Braga_CV.docx")
print("ok")
